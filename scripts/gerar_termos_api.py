"""
Script utilitário para gerar os arquivos oficiais em .docx e .pdf:
1. OFICIO_CONCESSAO_ACESSO_API.docx e .pdf
2. TERMO_DE_COMPROMISSO_E_SIGILO_LGPD.docx e .pdf
"""
import os
import sys
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib import colors as rl_colors

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_DIR = os.path.join(BASE_DIR, 'docs', 'modelos_oficiais')
os.makedirs(OUTPUT_DIR, exist_ok=True)

def gerar_docx_termo():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Cabecalho
    p_cab = doc.add_paragraph()
    p_cab.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_cab.add_run("ESTADO DO PARÁ\nPREFEITURA MUNICIPAL DE TOMÉ-AÇU\nSECRETARIA MUNICIPAL DE TRABALHO E ASSISTÊNCIA SOCIAL – SETAS\nCOORDENAÇÃO DO CADASTRO ÚNICO E PROGRAMA BOLSA FAMÍLIA\n")
    r1.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = RGBColor(30, 41, 59)

    p_tit = doc.add_paragraph()
    p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tit = p_tit.add_run("\nTERMO DE COMPROMISSO, CONFIDENCIALIDADE E RESPONSABILIDADE LGPD\n(INTEGRAÇÃO E ACESSO À API DO CADASTRO ÚNICO)\n")
    r_tit.bold = True
    r_tit.font.size = Pt(12)
    r_tit.font.color.rgb = RGBColor(4, 120, 87)

    texto_intro = (
        "Pelo presente instrumento, a entidade/setor solicitante abaixo qualificado formaliza o compromisso de "
        "confidencialidade, sigilo e conformidade no tratamento de dados pessoais custodiados pelo Sistema de Gestão do "
        "Cadastro Único e Visitas Domiciliares de Tomé-Açu/PA, em estrita observância à Lei Geral de Proteção de Dados (Lei nº 13.709/2018)."
    )
    p_intro = doc.add_paragraph(texto_intro)
    p_intro.paragraph_format.space_after = Pt(8)

    # Tabela de Dados Solicitante
    table = doc.add_table(rows=7, cols=2)
    table.style = 'Table Grid'
    dados = [
        ("Órgão / Entidade Solicitante:", "[Nome da Instituição / Secretaria]"),
        ("CNPJ / Matrícula Institucional:", "[00.000.000/0001-00]"),
        ("Responsável Legal do Órgão:", "[Nome do Secretário / Gestor]"),
        ("CPF do Responsável Legal:", "[000.000.000-00]"),
        ("Responsável Técnico / Desenvolvedor:", "[Nome do Técnico / TI]"),
        ("E-mail Institucional:", "[contato@instituicao.gov.br]"),
        ("Finalidade Específica da Integração:", "[Descrever a finalidade pactuada]")
    ]
    for i, (rotulo, valor) in enumerate(dados):
        row = table.rows[i]
        r_cell0 = row.cells[0].paragraphs[0].add_run(rotulo)
        r_cell0.bold = True
        r_cell0.font.size = Pt(9.5)
        r_cell1 = row.cells[1].paragraphs[0].add_run(valor)
        r_cell1.font.size = Pt(9.5)

    doc.add_paragraph()

    clausulas = [
        ("CLÁUSULA PRIMEIRA – DO OBJETO", "O presente Termo tem por objeto a concessão de Chave de API e credenciais técnicas para acesso controlado às rotinas do Sistema de Gestão do Cadastro Único da SETAS Tomé-Açu/PA."),
        ("CLÁUSULA SEGUNDA – DO SIGILO E NÃO COMPARTILHAMENTO", "O Solicitante compromete-se a manter sigilo absoluto sobre todas as informações socioeconômicas e credenciais recebidas, sendo terminantemente proibido o repasse, cessão ou compartilhamento da chave de API com terceiros."),
        ("CLÁUSULA TERCEIRA – DA CONFORMIDADE COM A LGPD", "O tratamento de dados pessoais dar-se-á estritamente para a finalidade pública justificada. O Solicitante obriga-se a notificar a SETAS em até 24 horas caso ocorra qualquer suspeita de vazamento ou incidente de segurança."),
        ("CLÁUSULA QUARTA – DAS PENALIDADES", "O descumprimento sujeitará o infrator e a entidade às sanções civis, administrativas e penais da Lei nº 13.709/2018 (LGPD) e da Lei de Improbidade Administrativa, além do cancelamento imediato do acesso."),
        ("CLÁUSULA QUINTA – DA VIGÊNCIA E FORO", "Vigência de 12 (doze) meses a contar da assinatura. Foro eleito: Comarca de Tomé-Açu/PA.")
    ]

    for tit, corpo in clausulas:
        p_c = doc.add_paragraph()
        r_c_tit = p_c.add_run(tit + "\n")
        r_c_tit.bold = True
        r_c_tit.font.size = Pt(10)
        r_c_corp = p_c.add_run(corpo)
        r_c_corp.font.size = Pt(9.5)
        p_c.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(f"\nTomé-Açu – PA, _____ de ____________________ de 20___.\n\n")

    p_ass = doc.add_paragraph()
    p_ass.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ass.add_run("_____________________________________________\nRESPONSÁVEL LEGAL DO ÓRGÃO SOLICITANTE\n\n\n")
    p_ass.add_run("_____________________________________________\nSECRETARIA MUNICIPAL DE TRABALHO E ASSISTÊNCIA SOCIAL (SETAS)\nPrefeitura Municipal de Tomé-Açu")

    docx_path = os.path.join(OUTPUT_DIR, 'TERMO_DE_COMPROMISSO_E_SIGILO_LGPD.docx')
    doc.save(docx_path)
    print(f"Gerado com sucesso: {docx_path}")

def gerar_docx_oficio():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Cabecalho
    p_cab = doc.add_paragraph()
    p_cab.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_cab.add_run("ESTADO DO PARÁ\nPREFEITURA MUNICIPAL DE TOMÉ-AÇU\nSECRETARIA MUNICIPAL DE TRABALHO E ASSISTÊNCIA SOCIAL – SETAS\n")
    r1.bold = True
    r1.font.size = Pt(11)
    r1.font.color.rgb = RGBColor(30, 41, 59)

    p_num = doc.add_paragraph()
    p_num.add_run("\nOFÍCIO Nº ______/20___ – SETAS/GAB\nTomé-Açu – PA, _____ de ____________________ de 20___.\n").bold = True

    p_dest = doc.add_paragraph()
    p_dest.add_run("A Sua Senhoria / Ilustríssimo(a) Senhor(a):\n").bold = True
    p_dest.add_run("[Nome do Gestor / Destinatário]\nCargo: [Cargo do Gestor]\nÓrgão/Secretaria: [Nome da Instituição Solicitante]\n")

    p_assunto = doc.add_paragraph()
    p_assunto.add_run("Assunto: ").bold = True
    p_assunto.add_run("Concessão de Acesso Técnico e Credenciais de Integração à API do Cadastro Único.\n")

    corpo = (
        "Senhor(a) Gestor(a),\n\n"
        "1. Cumprimentando-o(a) cordialmente, servimo-nos do presente para comunicar a aprovação e liberação do acesso técnico "
        "via API ao Sistema de Gestão do Cadastro Único e Visitas Domiciliares, em atendimento à solicitação formulada por essa instituição.\n\n"
        "2. A concessão do acesso tem por finalidade exclusiva subsidiar as ações de: [Descrever a finalidade pactuada].\n\n"
        "3. Os parâmetros oficiais de integração são:\n"
        "   • Ambiente: Produção\n"
        "   • URL Base da API: https://cadunico-sistema.onrender.com\n"
        "   • Tipo de Autenticação: Token Secreto (X-CRON-SECRET / Bearer Token)\n"
        "   • Credencial de Acesso: Enviada separadamente via cofre temporário criptografado ao responsável técnico indicado.\n\n"
        "4. Ressaltamos que a manutenção do acesso está condicionada ao estrito cumprimento das cláusulas do Termo de Compromisso e Sigilo LGPD "
        "(Lei Federal nº 13.709/2018) previamente firmado, sendo a guarda das credenciais de responsabilidade indelegável dessa instituição.\n\n"
        "Colocamo-nos à disposição para esclarecimentos adicionais.\n\n"
        "Respeitosamente,\n\n\n"
    )
    p_corpo = doc.add_paragraph(corpo)

    p_ass = doc.add_paragraph()
    p_ass.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ass.add_run("_____________________________________________\nSECRETÁRIO(A) MUNICIPAL DE TRABALHO E ASSISTÊNCIA SOCIAL\nSETAS · Prefeitura Municipal de Tomé-Açu")

    docx_path = os.path.join(OUTPUT_DIR, 'OFICIO_CONCESSAO_ACESSO_API.docx')
    doc.save(docx_path)
    print(f"Gerado com sucesso: {docx_path}")

def gerar_pdf_termo():
    pdf_path = os.path.join(OUTPUT_DIR, 'TERMO_DE_COMPROMISSO_E_SIGILO_LGPD.pdf')
    doc = SimpleDocTemplate(pdf_path, pagesize=A4, rightMargin=2*cm, leftMargin=2*cm, topMargin=1.5*cm, bottomMargin=1.5*cm)
    styles = getSampleStyleSheet()

    style_cab = ParagraphStyle('Cab', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=9, leading=12, alignment=1, textColor=rl_colors.HexColor('#1E293B'))
    style_tit = ParagraphStyle('Tit', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=11, leading=14, alignment=1, textColor=rl_colors.HexColor('#047857'))
    style_corp = ParagraphStyle('Corp', parent=styles['Normal'], fontName='Helvetica', fontSize=8.5, leading=11, textColor=rl_colors.HexColor('#1E293B'))
    style_claus = ParagraphStyle('Claus', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8.5, leading=11, textColor=rl_colors.HexColor('#0F172A'))

    story = [
        Paragraph("ESTADO DO PARÁ<br/>PREFEITURA MUNICIPAL DE TOMÉ-AÇU<br/>SECRETARIA MUNICIPAL DE TRABALHO E ASSISTÊNCIA SOCIAL – SETAS<br/>COORDENAÇÃO DO CADASTRO ÚNICO E PROGRAMA BOLSA FAMÍLIA", style_cab),
        Spacer(1, 0.2*cm),
        HRFlowable(width="100%", thickness=1.5, color=rl_colors.HexColor('#047857')),
        Spacer(1, 0.2*cm),
        Paragraph("TERMO DE COMPROMISSO, CONFIDENCIALIDADE E RESPONSABILIDADE LGPD<br/><font size=8>(INTEGRAÇÃO E ACESSO À API DO CADASTRO ÚNICO)</font>", style_tit),
        Spacer(1, 0.2*cm),
        Paragraph("Pelo presente instrumento, a entidade/setor solicitante abaixo qualificado formaliza o compromisso de confidencialidade, sigilo e conformidade no tratamento de dados pessoais custodiados pelo Sistema de Gestão do Cadastro Único e Visitas Domiciliares de Tomé-Açu/PA, nos termos da Lei Federal nº 13.709/2018 (LGPD).", style_corp),
        Spacer(1, 0.2*cm)
    ]

    t_data = [
        [Paragraph("<b>Órgão / Entidade Solicitante:</b>", style_corp), Paragraph("[Nome da Instituição Solicitante]", style_corp)],
        [Paragraph("<b>CNPJ / Matrícula:</b>", style_corp), Paragraph("[00.000.000/0001-00]", style_corp)],
        [Paragraph("<b>Responsável Legal / Titular:</b>", style_corp), Paragraph("[Nome do Secretário / Gestor]", style_corp)],
        [Paragraph("<b>Responsável Técnico (TI):</b>", style_corp), Paragraph("[Nome do Desenvolvedor / TI]", style_corp)],
        [Paragraph("<b>E-mail Institucional:</b>", style_corp), Paragraph("[contato@orgao.gov.br]", style_corp)],
        [Paragraph("<b>Finalidade da Integração:</b>", style_corp), Paragraph("[Finalidade pactuada]", style_corp)]
    ]
    t = Table(t_data, colWidths=[5.5*cm, 11*cm])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), rl_colors.HexColor('#F8FAFC')),
        ('BOX', (0,0), (-1,-1), 0.5, rl_colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, rl_colors.HexColor('#E2E8F0')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(t)
    story.append(Spacer(1, 0.25*cm))

    clausulas = [
        ("CLÁUSULA 1ª – OBJETO:", "Disponibilização de Chave de Autenticação Técnica (API Token) para consumo controlado do Sistema de Gestão do Cadastro Único da SETAS Tomé-Açu/PA."),
        ("CLÁUSULA 2ª – DO SIGILO:", "O Solicitante manterá sigilo absoluto sobre os dados socioeconômicos e credenciais recebidas, vedado o repasse, cessão ou compartilhamento da chave de API a terceiros."),
        ("CLÁUSULA 3ª – CONFORMIDADE LGPD:", "Tratamento estrito para a finalidade pública justificada. Obrigação de notificar a SETAS em até 24h sobre qualquer incidente de segurança ou suspeita de vazamento."),
        ("CLÁUSULA 4ª – PENALIDADES:", "O descumprimento sujeitará o infrator às sanções civis, administrativas e penais da Lei nº 13.709/2018 (LGPD), além da revogação imediata do acesso."),
        ("CLÁUSULA 5ª – VIGÊNCIA E FORO:", "Vigência de 12 meses. Foro eleito: Comarca de Tomé-Açu/PA.")
    ]
    for tit, corpo in clausulas:
        story.append(Paragraph(f"<b>{tit}</b> {corpo}", style_corp))
        story.append(Spacer(1, 0.15*cm))

    story.append(Spacer(1, 0.3*cm))
    story.append(Paragraph("Tomé-Açu – PA, _____ de ____________________ de 20___.", style_corp))
    story.append(Spacer(1, 0.6*cm))

    ass_data = [
        [Paragraph("____________________________________________<br/><b>RESPONSÁVEL LEGAL DO SOLICITANTE</b>", style_cab),
         Paragraph("____________________________________________<br/><b>SECRETÁRIO(A) MUNICIPAL – SETAS</b>", style_cab)]
    ]
    t_ass = Table(ass_data, colWidths=[8.5*cm, 8.5*cm])
    t_ass.setStyle(TableStyle([('ALIGN', (0,0), (-1,-1), 'CENTER')]))
    story.append(t_ass)

    doc.build(story)
    print(f"Gerado com sucesso: {pdf_path}")

def gerar_pdf_oficio():
    pdf_path = os.path.join(OUTPUT_DIR, 'OFICIO_CONCESSAO_ACESSO_API.pdf')
    doc = SimpleDocTemplate(pdf_path, pagesize=A4, rightMargin=2*cm, leftMargin=2*cm, topMargin=1.5*cm, bottomMargin=1.5*cm)
    styles = getSampleStyleSheet()

    style_cab = ParagraphStyle('Cab', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10, leading=13, alignment=1, textColor=rl_colors.HexColor('#1E293B'))
    style_corp = ParagraphStyle('Corp', parent=styles['Normal'], fontName='Helvetica', fontSize=9.5, leading=13, textColor=rl_colors.HexColor('#1E293B'))

    story = [
        Paragraph("ESTADO DO PARÁ<br/>PREFEITURA MUNICIPAL DE TOMÉ-AÇU<br/>SECRETARIA MUNICIPAL DE TRABALHO E ASSISTÊNCIA SOCIAL – SETAS", style_cab),
        Spacer(1, 0.2*cm),
        HRFlowable(width="100%", thickness=1.5, color=rl_colors.HexColor('#047857')),
        Spacer(1, 0.4*cm),
        Paragraph("<b>OFÍCIO Nº ______/20___ – SETAS/GAB</b><br/>Tomé-Açu – PA, _____ de ____________________ de 20___.", style_corp),
        Spacer(1, 0.4*cm),
        Paragraph("<b>A Sua Senhoria / Ilustríssimo(a) Senhor(a):</b><br/>[Nome do Gestor / Destinatário]<br/>Cargo: [Cargo do Gestor]<br/>Órgão/Secretaria: [Nome da Instituição Solicitante]", style_corp),
        Spacer(1, 0.4*cm),
        Paragraph("<b>Assunto:</b> Concessão de Acesso Técnico e Credenciais de Integração à API do Cadastro Único.", style_corp),
        Spacer(1, 0.4*cm),
        Paragraph("Senhor(a) Gestor(a),<br/><br/>"
                  "1. Cumprimentando-o(a) cordialmente, servimo-nos do presente para comunicar a <b>aprovação e liberação do acesso técnico</b> via API ao <b>Sistema de Gestão do Cadastro Único e Visitas Domiciliares</b>, em atendimento à solicitação formulada por essa entidade.<br/><br/>"
                  "2. A concessão do acesso tem por finalidade exclusiva subsidiar as ações pactuadas no requerimento oficial.<br/><br/>"
                  "3. Os parâmetros oficiais de conexão são:<br/>"
                  "&nbsp;&nbsp;&nbsp;• <b>Ambiente:</b> Produção<br/>"
                  "&nbsp;&nbsp;&nbsp;• <b>URL Base da API:</b> https://cadunico-sistema.onrender.com<br/>"
                  "&nbsp;&nbsp;&nbsp;• <b>Tipo de Autenticação:</b> Token Secreto (X-CRON-SECRET / Bearer Token)<br/>"
                  "&nbsp;&nbsp;&nbsp;• <b>Credencial de Acesso:</b> Enviada separadamente via cofre temporário seguro ao responsável técnico indicado.<br/><br/>"
                  "4. Ressaltamos que a manutenção do acesso está condicionada ao estrito cumprimento das cláusulas do Termo de Compromisso e Sigilo LGPD (Lei Federal nº 13.709/2018) previamente firmado.<br/><br/>"
                  "Colocamo-nos à disposição para esclarecimentos adicionais.<br/><br/>"
                  "Respeitosamente,", style_corp),
        Spacer(1, 1.2*cm),
        Paragraph("____________________________________________________<br/><b>SECRETÁRIO(A) MUNICIPAL DE TRABALHO E ASSISTÊNCIA SOCIAL</b><br/>SETAS · Prefeitura Municipal de Tomé-Açu", style_cab)
    ]
    doc.build(story)
    print(f"Gerado com sucesso: {pdf_path}")

if __name__ == '__main__':
    gerar_docx_termo()
    gerar_docx_oficio()
    gerar_pdf_termo()
    gerar_pdf_oficio()
    print("TODOS OS DOCUMENTOS FORAM GERADOS COM SUCESSO!")
